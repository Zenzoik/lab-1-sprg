#!/usr/bin/env python3
"""Build LR2_report.docx from docs/lab2-report.md in the layout of the LR1 report."""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1]
IMGDIR = sys.argv[2]
OUT = sys.argv[3]
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)


def run(p, text, bold=False, italic=False, size=None, mono=False, color=None, underline=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.underline = underline
    if mono:
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    else:
        r.font.name = FONT
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def add_hyperlink(p, url, text, size=None):
    part = p.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def inline(p, text, size=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run(p, part[2:-2].replace("`", ""), bold=True, size=size)
        elif part.startswith("`") and part.endswith("`"):
            run(p, part[1:-1], mono=True, size=(size or 14) - 2)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            add_hyperlink(p, m.group(2), m.group(1), size=size)
        else:
            run(p, part, size=size)


def para(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=None, after=6, keep=False, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if bold:
        run(p, text, bold=True, size=size)
    elif text:
        inline(p, text, size)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(rows):
    header = rows[0]
    body = rows[1:]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    usable = 17.0
    t.autofit = False
    widths = [max(len(c), 6) for c in header]
    for r_ in body:
        for i, c in enumerate(r_[:ncol]):
            widths[i] = max(widths[i], min(len(c), 60))
    tot = float(sum(widths))
    fs = 9.5 if ncol >= 6 else 10.5
    colw = [Cm(max(1.8, usable * w_ / tot)) for w_ in widths]
    scale = usable / sum(c.cm for c in colw)
    colw = [Cm(c.cm * scale) for c in colw]
    for ci, col in enumerate(t.columns):
        col.width = colw[ci]
    for ri, r_ in enumerate([header] + body):
        for ci in range(ncol):
            txt = r_[ci] if ci < len(r_) else ""
            cell = t.cell(ri, ci)
            cell.width = colw[ci]
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if ri == 0:
                run(cp, txt, bold=True, size=fs)
                shade(cell, "D9E2F3")
            else:
                inline(cp, txt, fs)
    para("", after=4)


def code_block(lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(0)
        run(p, ln if ln else " ", mono=True, size=9)
    para("", after=4)


fig_no = [0]


def figure(alt, path):
    full = os.path.join(IMGDIR, os.path.basename(path))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(full, width=Cm(16.0))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    run(c, alt)


# ---- title page (same layout as LR1) ----
for line in ["Міністерство освіти і науки України", "Національний технічний університет України",
             "«Київський політехнічний інститут імені Ігоря Сікорського»",
             "Факультет інформатики та обчислювальної техніки", "Кафедра обчислювальної техніки"]:
    para(line, WD_ALIGN_PARAGRAPH.CENTER, after=0)
for _ in range(5):
    para("", after=0)
para("Лабораторна робота № 2", WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=0)
para("з дисципліни", WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("«Системне програмування»", WD_ALIGN_PARAGRAPH.CENTER, after=0)
for _ in range(5):
    para("", after=0)
for line in ["ВИКОНАВ:", "студент IV курсу ФІОТ", "групи ІО-34", "Зернишкін Ілля Денисович", "Залікова № 3404"]:
    para(line, WD_ALIGN_PARAGRAPH.LEFT, after=0, indent=9.0)
para("", after=0)
for line in ["ПЕРЕВІРИВ:", "Гордієнко Н. Ю."]:
    para(line, WD_ALIGN_PARAGRAPH.LEFT, after=0, indent=9.0)
for _ in range(6):
    para("", after=0)
para("Київ – 2026", WD_ALIGN_PARAGRAPH.CENTER, after=0)
doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

# ---- body ----
lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
first_h2 = True
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        j = i + 1
        buf = []
        while not lines[j].startswith("```"):
            buf.append(lines[j])
            j += 1
        code_block(buf)
        i = j + 1
        continue
    if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-|]+\|$", lines[i + 1]):
        rows = []
        j = i
        while j < len(lines) and lines[j].startswith("|"):
            if not re.match(r"^\|[\s\-|]+\|$", lines[j]):
                cells = [c.strip() for c in lines[j].strip().strip("|").split(" | ")]
                if len(cells) == 1:
                    cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                rows.append(cells)
            j += 1
        table(rows)
        i = j
        continue
    m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", ln)
    if m:
        figure(m.group(1), m.group(2))
        i += 1
        continue
    if ln.startswith("## "):
        para(ln[3:], WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=8, keep=True)
        i += 1
        continue
    if ln.startswith("### "):
        para(ln[4:], WD_ALIGN_PARAGRAPH.LEFT, bold=True, after=6, keep=True)
        i += 1
        continue
    if re.match(r"^\d+\. ", ln):
        p = para("", after=3, indent=0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        inline(p, ln)
        i += 1
        continue
    if ln.startswith("- "):
        p = para("", after=3, indent=0.9)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        run(p, "• ")
        inline(p, ln[2:])
        i += 1
        continue
    if ln.strip() == "":
        i += 1
        continue
    para(ln)
    i += 1

# page numbers (footer, centered)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if tag:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tag)
        r._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = txt
        r._r.append(it)

doc.save(OUT)
print("saved", OUT)
